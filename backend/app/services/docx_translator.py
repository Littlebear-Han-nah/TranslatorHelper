import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from pathlib import Path
from typing import Dict, Any, Optional, Callable
from .qwen_client import QwenClient
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from ..utils.font_helper import get_chinese_font_name
import fitz

class DocxTranslator:
    """
    Word (.docx) 课件与论文翻译引擎
    支持段落提取、大模型翻译，并生成：
    1. 左右双栏对照排版的 Word 文档；
    2. 左右对照的双语 PDF 文档供直接下载与双栏预览。
    """

    def __init__(self, qwen_client: QwenClient):
        """
        初始化 Word 翻译器
        :param qwen_client: 通义千问客户端实例
        """
        self.client = qwen_client

    async def translate_docx(
        self,
        input_docx_path: str,
        output_dir: str,
        task_id: str,
        model: str = "qwen3.8-flash",
        mode: str = "academic",
        progress_callback: Optional[Callable[[Dict[str, Any]], None]] = None
    ) -> Dict[str, Any]:
        """
        执行 Word 文档翻译并生成双语对照 Word 与 PDF
        """
        out_path = Path(output_dir)
        preview_dir = out_path / "previews" / task_id
        preview_dir.mkdir(parents=True, exist_ok=True)

        doc = docx.Document(input_docx_path)
        
        # 提取非空段落
        paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        total_p = len(paragraphs)
        if total_p == 0:
            total_p = 1
            paragraphs = ["（文档中无文本段落）"]

        translated_paragraphs = []
        for idx, p_text in enumerate(paragraphs):
            if progress_callback:
                await progress_callback({
                    "task_id": task_id,
                    "stage": "translating",
                    "current_page": idx + 1,
                    "total_pages": total_p,
                    "percent": int(((idx) / total_p) * 80),
                    "message": f"正在翻译第 {idx + 1}/{total_p} 段落..."
                })

            trans = await self.client.translate_text(
                text=p_text,
                model=model,
                mode=mode
            )
            translated_paragraphs.append((p_text, trans))

        # 1. 生成双语对照 Word 文档 (采用双列表格对照布局：左列英文，右列中文)
        bilingual_docx = docx.Document()
        bilingual_docx.add_heading("文档双语对照翻译 (Bilingual Translation)", level=1)
        
        # 添加对照表格
        table = bilingual_docx.add_table(rows=1, cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        hdr_cells = table.rows[0].cells
        hdr_cells[0].text = "原文 (Original)"
        hdr_cells[1].text = "中文翻译 (Chinese Translation)"
        
        # 加粗表头
        for cell in hdr_cells:
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.bold = True
                    run.font.color.rgb = RGBColor(37, 99, 235)

        for orig, trans in translated_paragraphs:
            row_cells = table.add_row().cells
            row_cells[0].text = orig
            row_cells[1].text = trans

        bilingual_docx_path = str(out_path / f"{task_id}_bilingual.docx")
        bilingual_docx.save(bilingual_docx_path)

        # 2. 生成对应的左右对照 PDF
        if progress_callback:
            await progress_callback({
                "task_id": task_id,
                "stage": "rendering_pdf",
                "current_page": total_p,
                "total_pages": total_p,
                "percent": 90,
                "message": "正在生成双语对照 PDF 与高分预览..."
            })

        side_by_side_pdf_path = str(out_path / f"{task_id}_bilingual_side_by_side.pdf")
        self._render_bilingual_pdf(translated_paragraphs, side_by_side_pdf_path)

        # 3. 生成预览图
        pdf_doc = fitz.open(side_by_side_pdf_path)
        pages_preview = []
        for p_no in range(len(pdf_doc)):
            pix = pdf_doc[p_no].get_pixmap(dpi=150)
            img_name = f"page_{p_no+1}.png"
            img_path = str(preview_dir / img_name)
            pix.save(img_path)
            pages_preview.append({
                "page": p_no + 1,
                "orig_text": "",
                "trans_text": "",
                "orig_img": f"/outputs/previews/{task_id}/{img_name}",
                "trans_img": f"/outputs/previews/{task_id}/{img_name}",
            })
        pdf_doc.close()

        if progress_callback:
            await progress_callback({
                "task_id": task_id,
                "stage": "completed",
                "current_page": total_p,
                "total_pages": total_p,
                "percent": 100,
                "message": "Word 文档对照翻译完成！"
            })

        return {
            "task_id": task_id,
            "total_pages": len(pages_preview),
            "side_by_side_pdf": f"/outputs/{task_id}_bilingual_side_by_side.pdf",
            "bilingual_docx": f"/outputs/{task_id}_bilingual.docx",
            "pages": pages_preview,
        }

    def _render_bilingual_pdf(self, items, output_path: str):
        """
        将段落对生成左右对照排版的 A4 横版 PDF
        """
        font_name = get_chinese_font_name()
        # 使用 A4 横向尺寸 (842 x 595 pt)，非常适合左右对照
        page_w, page_h = 842, 595
        c = canvas.Canvas(output_path, pagesize=(page_w, page_h))

        margin = 40
        col_w = (page_w - margin * 2 - 20) / 2
        y = page_h - margin

        # 绘制页眉
        def draw_header():
            nonlocal y
            c.setFont(font_name, 12)
            c.setFillColor(colors.HexColor("#2563EB"))
            c.drawString(margin, page_h - 30, "原文 (Original)")
            c.drawString(margin + col_w + 20, page_h - 30, "中文对照 (Chinese Translation)")
            c.setStrokeColor(colors.HexColor("#CBD5E1"))
            c.setLineWidth(1)
            c.line(margin, page_h - 35, page_w - margin, page_h - 35)
            # 中线
            c.line(margin + col_w + 10, page_h - 35, margin + col_w + 10, margin)
            y = page_h - 50

        draw_header()

        for orig, trans in items:
            # 估算高度
            orig_lines = self._wrap_text(orig, col_w, 9.5, font_name, c)
            trans_lines = self._wrap_text(trans, col_w, 9.5, font_name, c)
            block_h = max(len(orig_lines), len(trans_lines)) * 14 + 10

            if y - block_h < margin:
                c.showPage()
                draw_header()

            # 绘制左侧英文
            c.setFont("Helvetica", 9)
            c.setFillColor(colors.HexColor("#1E293B"))
            cur_y = y
            for line in orig_lines:
                c.drawString(margin, cur_y, line)
                cur_y -= 14

            # 绘制右侧中文
            c.setFont(font_name, 9.5)
            c.setFillColor(colors.HexColor("#0F172A"))
            cur_y = y
            for line in trans_lines:
                c.drawString(margin + col_w + 20, cur_y, line)
                cur_y -= 14

            # 段落底部分割微线
            y -= block_h
            c.setStrokeColor(colors.HexColor("#F1F5F9"))
            c.setLineWidth(0.5)
            c.line(margin, y + 4, page_w - margin, y + 4)

        c.showPage()
        c.save()

    def _wrap_text(self, text, max_w, font_size, font_name, c):
        """
        折行辅助计算
        """
        chars = list(text)
        lines = []
        cur = ""
        for ch in chars:
            if c.stringWidth(cur + ch, font_name, font_size) <= max_w:
                cur += ch
            else:
                lines.append(cur)
                cur = ch
        if cur:
            lines.append(cur)
        return lines or [""]
