"""Generate small, synthetic, shareable fixtures for manual acceptance testing."""
from pathlib import Path
import io
import pymupdf as fitz
from PIL import Image
from docx import Document
from pptx import Presentation
from pptx.util import Inches
from openpyxl import Workbook

out = Path(__file__).resolve().parents[2] / 'artifacts' / 'fixtures'
out.mkdir(parents=True, exist_ok=True)
doc = fitz.open()
p = doc.new_page(width=595, height=842)
p.draw_rect(fitz.Rect(0,0,595,120), fill=(.88,.94,.9), color=None)
p.insert_text((42,62), 'Document Intelligence', fontsize=24, fontname='hebo', color=(.12,.3,.22))
p.insert_text((42,91), 'A synthetic fixture for document translation', fontsize=11)
p.insert_text((42,157), '1. Layout preservation', fontsize=16, fontname='hebo')
p.insert_textbox(fitz.Rect(42,178,283,268), 'Document translation preserves the original layout. Figures, equations and tables carry essential information. The translated page should remain readable.', fontsize=11, lineheight=1.5)
p.insert_textbox(fitz.Rect(313,178,553,268), 'A reliable system checks each text region before replacing the original. Difficult regions should be clearly marked for human review.', fontsize=11, lineheight=1.5)
p.insert_text((42,310),'2. Evaluation results', fontsize=16,fontname='hebo')
for y in (332,368,404,440): p.draw_line((42,y),(553,y),color=(.6,.7,.6))
for x in (42,300,553):p.draw_line((x,332),(x,440),color=(.6,.7,.6))
for x,y,t in [(52,354,'Model'),(310,354,'Accuracy'),(52,390,'Baseline'),(310,390,'82.5%'),(52,426,'Translation'),(310,426,'96.2%')]:p.insert_text((x,y),t,fontsize=11)
p.insert_text((42,485),'3. Figures and equations',fontsize=16,fontname='hebo')
p.insert_text((42,520),'E = mc2',fontsize=14)
image=Image.new('RGB',(240,110),'#d7e5d2')
for x in range(30,80):
 for y in range(40,100):image.putpixel((x,y),(60,104,74))
for x in range(110,160):
 for y in range(15,100):image.putpixel((x,y),(117,153,107))
buf=io.BytesIO();image.save(buf,format='PNG')
p.insert_image(fitz.Rect(313,500,553,610),stream=buf.getvalue())
p.insert_text((313,631),'Figure 1. A preserved image',fontsize=10)
p.insert_text((291,815),'1',fontsize=10)
doc.set_metadata({'title':'Document Intelligence - Synthetic Fixture','author':'TranslatorHelper Tests'})
doc.save(out/'layout.pdf')
p.get_pixmap(matrix=fitz.Matrix(2,2)).save(out/'scan.png')
doc.close()
Image.open(out/'scan.png').convert('RGB').save(out/'scan.jpg',quality=95)
word=Document();word.add_heading('Document Intelligence',0);word.add_paragraph('Document translation preserves the original layout.')
word.add_heading('Evaluation results',1);table=word.add_table(rows=2,cols=2);table.style='Table Grid';table.cell(0,0).text='Model';table.cell(0,1).text='Accuracy';table.cell(1,0).text='Baseline';table.cell(1,1).text='82.5%';word.save(out/'document.docx')
prs=Presentation();slide=prs.slides.add_slide(prs.slide_layouts[5]);slide.shapes.title.text='Document Intelligence';box=slide.shapes.add_textbox(Inches(1), Inches(2), Inches(8), Inches(2));box.text='Document translation preserves figures and equations.';prs.save(out/'slides.pptx')
wb=Workbook();ws=wb.active;ws.title='Results';ws['A1']='Evaluation results';ws.merge_cells('A1:C1');ws['A2']='Model';ws['B2']='Accuracy';ws['A3']='Baseline';ws['B3']=.825;ws['B3'].number_format='0.0%';ws['C3']='=B3*100';ws.column_dimensions['A'].width=26;ws.column_dimensions['B'].width=18;wb.save(out/'sheet.xlsx')
print('Synthetic PDF, PNG, JPG, DOCX, PPTX and XLSX fixtures generated.')
