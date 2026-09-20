import asyncio
import io
import json
import zipfile
from pathlib import Path
import httpx
import pytest
import pymupdf as fitz
from PIL import Image
from lxml import etree
from backend.app.services.model_adapter import CompatibleAdapter, ModelError
from backend.app.services.layout_engine import Block, extract_blocks, render_blocks, translate_pdf
from backend.app.services.office_engine import translate_native, translate_office, validate_office, W, A, S

def run(coroutine): return asyncio.run(coroutine)

class FakeAdapter:
    async def translate(self, blocks, *args):
        return {b['id']: {'Layout matters': '版面很重要', 'Hello world': '你好世界', 'Accuracy': '准确率', 'Model': '模型', 'Figure 1': '图 1'}.get(b['text'], '文档翻译保留原始结构') for b in blocks}

def noop(*args): pass

def test_adapter_validates_ids_and_protected_formulas():
    def respond(request):
        body=json.loads(request.content)
        items=json.loads(body['messages'][1]['content'])
        return httpx.Response(200,json={'choices':[{'message':{'content':json.dumps({item['id']:item['text'].replace('Equation', '公式') for item in items})}}]})
    adapter=CompatibleAdapter(api_key='test-only',transport=httpx.MockTransport(respond))
    value=run(adapter.translate([{'id':'a','text':'Equation $x^2$ [1]'}], 'custom-model'))
    assert value == {'a':'公式 $x^2$ [1]'}

@pytest.mark.parametrize('reply', ['{"wrong":"中文"}', '{"a":""}', '{"a":"公式"}', 'not json'])
def test_adapter_rejects_incomplete_response(reply):
    adapter=CompatibleAdapter(api_key='test-only',transport=httpx.MockTransport(lambda _: httpx.Response(200,json={'choices':[{'message':{'content':reply}}]})))
    with pytest.raises(ModelError): run(adapter.translate([{'id':'a','text':'Equation $x$'}], 'test'))

def test_adapter_does_not_leak_provider_error():
    adapter=CompatibleAdapter(api_key='test-only',transport=httpx.MockTransport(lambda _: httpx.Response(401,text='private details')))
    with pytest.raises(ModelError,match='拒绝访问'): run(adapter.chat('test',[]))


def test_adapter_recovers_when_large_json_batch_loses_ids():
    calls = []
    def respond(request):
        items = json.loads(json.loads(request.content)['messages'][1]['content'])
        calls.append(len(items))
        selected = items[:1] if len(items) > 1 else items
        answer = {item['id']: '中文' for item in selected}
        return httpx.Response(200, json={'choices': [{'message': {'content': json.dumps(answer)}}]})
    adapter = CompatibleAdapter(api_key='test-only', transport=httpx.MockTransport(respond))
    blocks = [{'id': str(i), 'text': f'English sentence {i}'} for i in range(5)]
    assert run(adapter.translate(blocks, 'test')) == {str(i): '中文' for i in range(5)}
    assert calls[0] == 5 and calls.count(1) == 5


def test_adapter_translates_repeated_slide_text_once():
    requested = []
    def respond(request):
        items = json.loads(json.loads(request.content)['messages'][1]['content'])
        requested.extend(items)
        return httpx.Response(200, json={'choices': [{'message': {
            'content': json.dumps({item['id']: '你好世界' for item in items})}}]})
    adapter = CompatibleAdapter(api_key='test-only', transport=httpx.MockTransport(respond))
    blocks = [{'id': f'slide-{i}', 'text': 'Hello world'} for i in range(50)]
    result = run(adapter.translate(blocks, 'test'))
    assert len(requested) == 1
    assert len(result) == 50
    assert set(result.values()) == {'你好世界'}

def test_layout_never_deletes_text_when_translation_overflows():
    doc=fitz.open(); page=doc.new_page()
    page.insert_text((40,50),'Hello world',fontsize=12)
    blocks=extract_blocks(page)
    blocks[0].translation='非常长的完整翻译' * 300
    render_blocks(page,page,blocks)
    assert blocks[0].status=='review'
    assert 'Hello world' in page.get_text()
    assert len(page.get_text()) < 100
    doc.close()

def test_pdf_preserves_images_drawings_dimensions_and_page_numbers(tmp_path):
    doc=fitz.open(); page=doc.new_page(width=500,height=700)
    page.draw_rect(fitz.Rect(30,30,470,90), fill=(.85,.92,.87),color=None)
    page.insert_text((40,60),'Layout matters',fontsize=20)
    page.insert_text((40,135),'Hello world',fontsize=12)
    page.insert_text((245,680),'1',fontsize=10)
    im=Image.new('RGB',(60,60),'#995522'); buf=io.BytesIO(); im.save(buf,format='PNG')
    page.insert_image(fitz.Rect(40,200,140,300),stream=buf.getvalue())
    page.draw_line((40,350),(450,350),color=(0,0,0))
    original=tmp_path/'input.pdf'; doc.save(original); doc.close()
    result=run(translate_pdf(original,tmp_path,FakeAdapter(),'test','academic','',noop,noop))
    with fitz.open(tmp_path/'translated.pdf') as translated, fitz.open(original) as source:
        assert len(translated)==len(source)==1
        assert translated[0].rect == source[0].rect
        assert len(translated[0].get_images()) == len(source[0].get_images())
        assert '1' in translated[0].get_text()
        assert '版面很重要' in translated[0].get_text()
        assert 'Layout matters' not in translated[0].get_text()
        clip=fitz.Rect(40,200,140,300)
        assert translated[0].get_pixmap(clip=clip).samples==source[0].get_pixmap(clip=clip).samples
        # Transparent redaction must preserve the colored heading background.
        clip=fitz.Rect(32,32,38,38)
        assert translated[0].get_pixmap(clip=clip).samples==source[0].get_pixmap(clip=clip).samples
    assert result['pages'][0]['blocks'][-1]['status'] == 'preserved'

def test_formula_font_is_protected():
    doc=fitz.open(); p=doc.new_page()
    p.insert_text((30,40),'sin(x)',fontname='symb')
    blocks=extract_blocks(p)
    assert blocks[0].kind=='protected'
    doc.close()

def test_multi_page_courseware_exports_every_bilingual_page(tmp_path):
    # With 48 pages, the former interleaved edit/graft loop crashed on page 2
    # at 16% with FzErrorArgument: source object number out of range.
    source = tmp_path / 'courseware.pdf'
    with fitz.open() as doc:
        for index in range(48):
            page = doc.new_page(width=720, height=405)
            page.draw_rect(fitz.Rect(20, 20, 700, 90), fill=(.85, .92, .87), color=None)
            page.insert_text((40, 60), 'Layout matters', fontsize=24)
            page.insert_text((40, 130), 'Hello world', fontsize=18)
            page.insert_text((350, 385), str(index+1), fontsize=10)
        doc.save(source)
    progress = []
    result = run(translate_pdf(source, tmp_path, FakeAdapter(), 'test', 'courseware', '',
                               lambda *args: progress.append(args), noop))
    with fitz.open(tmp_path/'translated.pdf') as translated, fitz.open(tmp_path/'bilingual.pdf') as paired:
        assert len(result['pages']) == len(translated) == len(paired) == 48
        for index in range(48):
            assert '版面很重要' in translated[index].get_text()
            assert 'Layout matters' not in translated[index].get_text()
            left = paired[index].get_text(clip=fitz.Rect(0, 0, 720, 405))
            right = paired[index].get_text(clip=fitz.Rect(720, 0, 1440, 405))
            assert 'Layout matters' in left and '版面很重要' not in left
            assert '版面很重要' in right and 'Layout matters' not in right
            assert str(index+1) in left.splitlines() and str(index+1) in right.splitlines()
            # Exported vector pages must render, not merely contain text objects.
            assert paired[index].get_pixmap(matrix=fitz.Matrix(.25, .25)).width == 360
    assert all(a[1] <= b[1] for a, b in zip(progress, progress[1:]))

def test_bilingual_export_keeps_blank_slides_and_supports_cancellation(tmp_path, monkeypatch):
    from backend.app.services import layout_engine
    monkeypatch.setattr(layout_engine, 'add_ocr', lambda page, blocks: (blocks, []))
    source = tmp_path/'blank.pdf'
    with fitz.open() as doc:
        doc.new_page()
        doc.new_page().insert_text((40, 50), 'Hello world')
        doc.save(source)
    run(translate_pdf(source, tmp_path, FakeAdapter(), 'test', 'courseware', '', noop, noop))
    with fitz.open(tmp_path/'bilingual.pdf') as doc:
        assert len(doc) == 2
        assert not doc[0].get_text().strip()
        assert '你好世界' in doc[1].get_text()
    class Cancelled(Exception): pass
    assembling = False
    def progress(stage, percent, message):
        nonlocal assembling
        if '双栏 PDF' in message:
            assembling = True
    def cancel():
        if assembling:
            raise Cancelled()
    with pytest.raises(Cancelled):
        run(translate_pdf(source, tmp_path, FakeAdapter(), 'test', 'courseware', '', progress, cancel))

def test_word_preserves_tables_equations_styles_and_media(tmp_path):
    from docx import Document
    from docx.oxml import OxmlElement
    doc=Document(); p=doc.add_paragraph(); run_node=p.add_run('Hello world'); run_node.bold=True
    math=OxmlElement('m:oMath'); math_run=OxmlElement('m:r'); math_text=OxmlElement('m:t'); math_text.text='x² + y²'; math_run.append(math_text); math.append(math_run); p._p.append(math)
    doc.add_table(rows=1,cols=1).cell(0,0).text='Accuracy'
    image=tmp_path/'image.png'; Image.new('RGB',(20,20),'red').save(image); doc.add_picture(str(image))
    source=tmp_path/'input.docx'; target=tmp_path/'translated.docx'; doc.save(source)
    run(translate_native(source,target,FakeAdapter(),'test','academic','',noop,noop))
    with zipfile.ZipFile(source) as src,zipfile.ZipFile(target) as dst:
        assert src.namelist()==dst.namelist()
        for name in src.namelist():
            if name.startswith('word/media/') or name.endswith('.rels'): assert src.read(name)==dst.read(name)
        tree=etree.fromstring(dst.read('word/document.xml'))
        assert 'x² + y²' in ''.join(tree.itertext())
        assert tree.find('.//{'+W+'}b') is not None
        assert tree.find('.//{'+W+'}tbl') is not None
        assert '你好世界' in ''.join(tree.itertext())

def test_powerpoint_preserves_slide_geometry_and_picture(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches
    prs=Presentation(); slide=prs.slides.add_slide(prs.slide_layouts[6])
    box=slide.shapes.add_textbox(Inches(1),Inches(1),Inches(5),Inches(1)); box.text='Hello world'
    image=tmp_path/'pic.png'; Image.new('RGB',(20,20),'green').save(image); slide.shapes.add_picture(str(image), Inches(2), Inches(3))
    source=tmp_path/'in.pptx'; target=tmp_path/'out.pptx'; prs.save(source)
    run(translate_native(source,target,FakeAdapter(),'test','academic','',noop,noop))
    translated=Presentation(target)
    assert len(translated.slides)==1
    assert translated.slides[0].shapes[0].text=='你好世界'
    assert translated.slides[0].shapes[0].left==box.left
    with zipfile.ZipFile(source) as src, zipfile.ZipFile(target) as dst:
        for name in src.namelist():
            if name.startswith('ppt/media/'): assert src.read(name)==dst.read(name)


def test_large_powerpoint_uses_translated_slides_for_pdf_without_second_model_pass(tmp_path, monkeypatch):
    from pptx import Presentation
    from pptx.util import Inches
    from backend.app.services import office_engine
    deck = Presentation()
    for index in range(50):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1)).text = 'Hello world'
    source = tmp_path / 'deck.pptx'
    deck.save(source)

    def fake_to_pdf(path, output):
        result = Path(output) / f'{Path(path).stem}.pdf'
        translated = Path(path).stem == 'translated'
        with fitz.open() as pdf:
            for _ in range(50):
                page = pdf.new_page(width=720, height=405)
                page.insert_text((40, 60), 'Translated' if translated else 'Hello world')
            pdf.save(result)
        return result

    class CountingAdapter(FakeAdapter):
        calls = 0
        async def translate(self, blocks, *args):
            self.calls += 1
            return await super().translate(blocks, *args)

    monkeypatch.setattr(office_engine, 'to_pdf', fake_to_pdf)
    adapter = CountingAdapter()
    result = run(translate_office(source, tmp_path, adapter, 'test', 'courseware', '', noop, noop))
    assert adapter.calls == 1
    assert len(result['pages']) == 50
    with fitz.open(tmp_path / 'bilingual.pdf') as paired:
        assert len(paired) == 50
        assert 'Hello world' in paired[49].get_text()
        assert 'Translated' in paired[49].get_text()


def test_powerpoint_chart_labels_are_translated(tmp_path):
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    chart = CategoryChartData()
    chart.categories = ['Revenue', 'Profit']
    chart.add_series('Sales', [2, 3])
    slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1),
                           Inches(5), Inches(3), chart)
    source, target = tmp_path / 'chart.pptx', tmp_path / 'chart-translated.pptx'
    deck.save(source)
    run(translate_native(source, target, FakeAdapter(), 'test', 'academic', '', noop, noop))
    with zipfile.ZipFile(target) as archive:
        chart_xml = archive.read('ppt/charts/chart1.xml').decode()
        assert 'Revenue' not in chart_xml
        assert 'Profit' not in chart_xml
        assert 'Sales' not in chart_xml
        assert '文档翻译保留原始结构' in chart_xml

def test_excel_preserves_formulas_merges_styles_and_formula_literals(tmp_path):
    from openpyxl import Workbook,load_workbook
    from openpyxl.styles import Font
    wb=Workbook(); ws=wb.active; ws['A1']='Hello world'; ws['A1'].font=Font(bold=True)
    ws.merge_cells('A1:B1'); ws['A2']=12; ws['B2']='=A2*2'; ws['C1']='Active'; ws['D1']='=COUNTIF(C1:C10,"Active")'
    source=tmp_path/'in.xlsx';target=tmp_path/'out.xlsx';wb.save(source)
    run(translate_native(source,target,FakeAdapter(),'test','academic','',noop,noop))
    result=load_workbook(target)
    assert result.active['A1'].value=='你好世界'
    assert result.active['A1'].font.bold
    assert str(result.active.merged_cells)=='A1:B1'
    assert result.active['A2'].value==12
    assert result.active['B2'].value=='=A2*2'
    assert result.active['C1'].value=='Active'
    assert result.active['D1'].value=='=COUNTIF(C1:C10,"Active")'

def test_office_rejects_external_fetch_relationship(tmp_path):
    path=tmp_path/'unsafe.docx'
    with zipfile.ZipFile(path,'w') as z:
        z.writestr('[Content_Types].xml','<Types/>')
        z.writestr('word/document.xml','<document/>')
        z.writestr('_rels/.rels','<Relationships><Relationship TargetMode="External" Type="externalLink" Target="https://example.invalid/private"/></Relationships>')
    with pytest.raises(ValueError,match='外部数据'):validate_office(path)

def test_table_cells_keep_alignment_and_never_translate_numbers():
    doc=fitz.open();p=doc.new_page()
    for y in [40,80,120]:p.draw_line((40,y),(340,y))
    for x in [40,190,340]:p.draw_line((x,40),(x,120))
    p.insert_text((50,65),'Model');p.insert_text((200,65),'Accuracy')
    p.insert_text((50,105),'Baseline');p.insert_text((200,105),'82.5%')
    blocks=extract_blocks(p)
    cells=[b for b in blocks if b.kind=='table']
    assert len(cells)==4
    assert next(b for b in cells if b.text=='82.5%').status=='preserved'
    assert next(b for b in cells if b.text=='Model').bbox[0]==50
    doc.close()

def test_paragraph_reflows_without_artificial_line_breaks():
    doc=fitz.open();p=doc.new_page()
    p.insert_textbox(fitz.Rect(40,40,180,180),'Document translation preserves the original layout and all essential information.',fontsize=12,lineheight=1.5)
    blocks=extract_blocks(p)
    assert len(blocks)==1
    assert '\n' not in blocks[0].text
    assert blocks[0].line_height >= 1.2
    doc.close()

def test_rotation_keeps_visual_page_dimensions(tmp_path):
    doc=fitz.open();p=doc.new_page(width=500,height=700);p.insert_text((40,60),'Hello world');p.set_rotation(90)
    source=tmp_path/'rotated.pdf';doc.save(source);doc.close()
    result=run(translate_pdf(source,tmp_path,FakeAdapter(),'test','academic','',noop,noop))
    assert result['pages'][0]['width']==700
    assert result['pages'][0]['height']==500

def test_two_column_paragraphs_merge_without_crossing_columns():
    doc=fitz.open();p=doc.new_page()
    p.insert_textbox(fitz.Rect(40,40,250,200),'Document translation preserves the original layout and all essential information in the left column.',fontsize=12,lineheight=1.5)
    p.insert_textbox(fitz.Rect(300,40,510,200),'A reliable document system keeps the right column separate and prevents mixed reading order.',fontsize=12,lineheight=1.5)
    blocks=extract_blocks(p)
    assert len(blocks)==2
    assert 'left column' in blocks[0].text
    assert 'right column' in blocks[1].text
    assert blocks[0].bbox[2] < blocks[1].bbox[0]
    doc.close()

def test_ocr_does_not_merge_table_columns(monkeypatch):
    import pytesseract
    from backend.app.services import layout_engine
    monkeypatch.setattr(layout_engine.shutil,'which',lambda _: '/test/tesseract')
    monkeypatch.setattr(pytesseract,'image_to_data',lambda *a,**k: {'text':['Model','Accuracy','82.5%'], 'block_num':[1,1,1], 'par_num':[1,1,1], 'line_num':[1,1,2], 'left':[20,250,250], 'top':[40,40,90], 'width':[60,80,60], 'height':[16,16,16], 'conf':[99,99,99]})
    doc=fitz.open();page=doc.new_page(width=400,height=200)
    blocks,_=layout_engine.add_ocr(page,[])
    assert [b.text for b in blocks]==['Model','Accuracy']
    assert blocks[0].bbox[2] < blocks[1].bbox[0]
    doc.close()


def test_ocr_checks_large_image_even_when_page_has_native_text(monkeypatch):
    import pytesseract
    from backend.app.services import layout_engine
    monkeypatch.setattr(layout_engine.shutil, 'which', lambda _: '/test/tesseract')
    monkeypatch.setattr(pytesseract, 'image_to_data', lambda *a, **k: {
        'text': ['Embedded', 'English'], 'block_num': [1, 1],
        'par_num': [1, 1], 'line_num': [1, 1], 'left': [200, 310],
        'top': [200, 200], 'width': [100, 90], 'height': [24, 24],
        'conf': [99, 99]})
    with fitz.open() as doc:
        page = doc.new_page(width=400, height=300)
        page.insert_text((10, 20), 'Native English heading with enough characters to trigger old skip')
        page.insert_text((10, 280), 'Additional native English text')
        image = Image.new('RGB', (300, 200), 'white')
        data = io.BytesIO(); image.save(data, format='PNG')
        page.insert_image(fitz.Rect(50, 50, 350, 250), stream=data.getvalue())
        blocks = extract_blocks(page)
        assert sum(len(block.text) for block in blocks) >= 50
        found, _ = layout_engine.add_ocr(page, blocks)
        assert any(block.kind == 'ocr' and block.text == 'Embedded English' for block in found)
